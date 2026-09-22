"""Text and metadata extraction for uploaded documents."""

import io
import re

import pymupdf
from docx import Document

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")
WORDS_PER_PAGE = 500
WORDS_PER_MINUTE = 230


class ExtractionError(Exception):
    pass


def _clean(text):
    # Re-join words hyphenated across line breaks, then collapse excess blank lines.
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _from_pdf(file_bytes):
    try:
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise ExtractionError(f"Could not open PDF: {e}")
    if doc.needs_pass:
        raise ExtractionError("This PDF is password-protected. Please upload an unlocked copy.")
    text = "\n".join(page.get_text() for page in doc)
    meta = doc.metadata or {}
    return text, {
        "title": (meta.get("title") or "").strip(),
        "authors": (meta.get("author") or "").strip(),
        "pages": doc.page_count,
    }


def _from_docx(file_bytes):
    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ExtractionError(f"Could not open DOCX: {e}")
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    props = doc.core_properties
    return "\n".join(parts), {
        "title": (props.title or "").strip(),
        "authors": (props.author or "").strip(),
        "pages": None,
    }


def _from_text(file_bytes):
    return file_bytes.decode("utf-8", errors="replace"), {"title": "", "authors": "", "pages": None}


def extract(filename, file_bytes):
    """Return (text, info) for a supported file, raising ExtractionError on failure."""
    name = filename.lower()
    if name.endswith(".pdf"):
        text, info = _from_pdf(file_bytes)
    elif name.endswith(".docx"):
        text, info = _from_docx(file_bytes)
    elif name.endswith((".txt", ".md")):
        text, info = _from_text(file_bytes)
    else:
        raise ExtractionError("Unsupported file type. Please upload a PDF, DOCX, or TXT file.")

    text = _clean(text)
    words = len(text.split())
    if info["pages"] is None:
        info["pages"] = max(1, round(words / WORDS_PER_PAGE))
    info["words"] = words
    info["reading_minutes"] = max(1, round(words / WORDS_PER_MINUTE))
    return text, info
